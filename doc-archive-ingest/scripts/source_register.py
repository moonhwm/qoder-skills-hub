#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
source_register.py —— 出处登记册生成器（读清单 JSON + 下载结果 → 出处登记册.md）
=============================================================================
输入：
  --manifest  文件清单 JSON（netdisk_browse.py 产物，可重复传入多个链接的清单）
  --evidence  下载尝试证据 JSON（可选；用于判定各文件下载成败与失败原因、轮次时间戳）
  --archive-root  归档根目录（可选；对已下载文件就地抽取摘要）
输出：出处登记册.md，逐文件字段：文件 | 大小 | 来源链接 | 获取时间 | 内容摘要 |
      版权与使用注记 | 轮次时间戳。字段规范详见 references/source_register_spec.md。

摘要抽取规则（已获取正文时）：
  .pdf  → 首页起前 500 字（pymupdf）；.docx → 正文段落前 500 字（python-docx）；
  .xlsx → 首个工作表表头行（openpyxl）；.txt/.md/.csv → 直接读前 500 字；
  图片 → 注明"图片文件，未做 OCR"；其余 → 依据文件名判断。
  依赖库缺失时如实注明"未能抽取（缺少 …）"，不编造摘要。
未获取正文时：诚实标注"未获取正文（原因；本管线不绕过该限制）"，依据文件名判断。

版权注记默认模板（字幕组式）：
  本文件由用户提供的公开分享链接获取，仅供用户个人学习使用，出处见链接；
  如涉第三方版权，权利归原权利人所有。

用法：
    python3 source_register.py --manifest 清单1.json [--manifest 清单2.json ...]
        [--evidence 下载尝试证据.json] [--archive-root 目录] [--out 出处登记册.md]
    python3 source_register.py --smoke    # 合成清单自测，不联网

依赖：标准库；摘要抽取按需惰性加载 pymupdf / python-docx / openpyxl。
"""
import argparse
import json
import os
import sys
import tempfile
from datetime import datetime, timedelta, timezone

COPYRIGHT_NOTE = ("本文件由用户提供的公开分享链接获取，仅供用户个人学习使用，"
                  "出处见链接；如涉第三方版权，权利归原权利人所有")
SUMMARY_LIMIT = 500  # 摘要字符上限
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff"}


def now_str():
    return datetime.now(timezone(timedelta(hours=8))).strftime("%Y-%m-%d %H:%M:%S +0800")


# ---------------------------------------------------------------- 摘要抽取
def _normalize(text):
    return " ".join(text.split())


def extract_summary(path):
    """对已下载到本地的文件抽取内容摘要。失败时返回诚实标注，不编造。"""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            try:
                import fitz
            except ImportError:
                return "PDF 文件；未能抽取正文（缺少 pymupdf）"
            doc = fitz.open(path)
            text = "".join(page.get_text("text") for page in doc)
            doc.close()
            text = _normalize(text)
            return (f"PDF 正文前{SUMMARY_LIMIT}字：{text[:SUMMARY_LIMIT]}"
                    if text else "PDF 无文本层（可能为扫描件），未做 OCR")
        if ext == ".docx":
            try:
                import docx
            except ImportError:
                return "DOCX 文件；未能抽取正文（缺少 python-docx）"
            d = docx.Document(path)
            text = _normalize("".join(p.text + "\n" for p in d.paragraphs))
            return (f"DOCX 正文前{SUMMARY_LIMIT}字：{text[:SUMMARY_LIMIT]}"
                    if text else "DOCX 正文为空")
        if ext == ".xlsx":
            try:
                import openpyxl
            except ImportError:
                return "XLSX 文件；未能抽取表头（缺少 openpyxl）"
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            ws = wb.worksheets[0] if wb.worksheets else None
            header = []
            if ws:
                for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
                    header = [str(c) if c is not None else "" for c in row]
                    break
            wb.close()
            header = [h for h in header if h]
            return (f"XLSX 首表表头：{' | '.join(header)}"
                    if header else "XLSX 首表无表头或为空表")
        if ext in (".txt", ".md", ".csv"):
            with open(path, encoding="utf-8", errors="replace") as f:
                text = _normalize(f.read(SUMMARY_LIMIT * 2))
            return f"文本前{SUMMARY_LIMIT}字：{text[:SUMMARY_LIMIT]}"
        if ext in IMAGE_EXTS:
            return f"图片文件（{ext[1:].upper()}），未做 OCR"
        return f"依据文件名判断：{os.path.splitext(os.path.basename(path))[0]}"
    except Exception as e:
        return f"摘要抽取失败（{type(e).__name__}: {e}），仅登记文件名"


def no_body_summary(name, reason):
    """未获取正文时的诚实标注格式。"""
    stem, ext = os.path.splitext(os.path.basename(name))
    ext_tag = ext[1:].upper() if ext else "未知类型"
    reason = reason or "未尝试下载"
    return (f"依据文件名判断：{stem}（{ext_tag}）；未获取正文（{reason}；"
            f"本管线不绕过该限制）")


# ---------------------------------------------------------------- 证据索引
def load_evidence_index(evidence_path):
    """把证据 JSON 拍平成 relPath -> {record, rounds:[(round, fetched_at)]}。"""
    index = {}
    if not evidence_path or not os.path.exists(evidence_path):
        return index
    with open(evidence_path, encoding="utf-8") as f:
        evidence = json.load(f)
    round_ts = {k: (v.get("fetched_at") or "")
                for k, v in evidence.items() if isinstance(v, dict)}
    for rkey in sorted(evidence):
        node = evidence[rkey]
        if not isinstance(node, dict):
            continue
        for rec in node.get("download_attempts", []):
            rel = rec.get("relPath")
            if not rel:
                continue
            slot = index.setdefault(rel, {"record": None, "rounds": []})
            slot["rounds"].append((rkey, round_ts.get(rkey, "")))
            slot["record"] = rec  # 后面的轮次覆盖前面的，取最新状态
    return index


def reason_of(rec):
    """把下载记录转成人类可读原因。"""
    if not rec:
        return "未在证据中找到下载记录"
    code = rec.get("errorCode")
    status = rec.get("http_status")
    if status == 403:
        return f"分享者设置了“仅预览/禁止下载”权限，下载接口返回 403 {code or ''}".strip()
    return f"下载未成功（HTTP {status} {code or ''}：{(rec.get('detailMsg') or '')[:80]}）".strip()


# ---------------------------------------------------------------- 登记册生成
def build_register(manifest_paths, evidence_path, archive_root, out_path):
    evidence_index = load_evidence_index(evidence_path)
    lines = ["# 出处登记册", ""]
    lines.append(f"- 登记时间：{now_str()}")
    lines.append("- 归档管线：坚果云 pubDIRBrowse/pubDIRLink 或百度网盘 xpan/分享页通道"
                 "枚举 → （授权范围内）下载 → 本登记册（不绕过任何权限设置，不伪造下载成功）")

    sections = []
    total_files = total_ok = total_fail = 0
    for mi, mpath in enumerate(manifest_paths, 1):
        with open(mpath, encoding="utf-8") as f:
            manifest = json.load(f)
        # 兼容坚果云清单（source）与百度网盘清单（source_url，baidu_share_browse.py 产物）
        share_url = manifest.get("source") or manifest.get("source_url", "")
        fetched_at = manifest.get("fetched_at", "")
        rnd = manifest.get("round")
        objects = manifest.get("objects", [])
        files = [o for o in objects if o.get("type") == "file"]
        dirs = [o for o in objects if o.get("type") == "directory"]
        total_files += len(files)

        rows, n_ok, n_fail = [], 0, 0
        for o in files:
            rel = o.get("relPath", "")
            name = o.get("name") or os.path.basename(rel)
            size = o.get("size")
            mtime = o.get("mtime")
            slot = evidence_index.get(rel, {})
            rec = slot.get("record")
            rounds = slot.get("rounds") or []
            round_ts = "；".join(f"{rk} {ts}".strip() for rk, ts in rounds) or "—"

            saved_rel = rec.get("saved_path") if rec else None
            local = None
            if saved_rel and archive_root:
                # 证据中的 saved_path 相对各链接子目录，拼上 link 键定位本地文件
                cand = os.path.join(archive_root, rec["link"], saved_rel)
                if os.path.exists(cand):
                    local = cand
            if local:
                summary = extract_summary(local)
                n_ok += 1
            else:
                summary = no_body_summary(name, reason_of(rec))
                n_fail += 1

            title = f"`{rel}`" + (f"（分享方修改于 {mtime}）" if mtime else "")
            rows.append("| %s | %s | %s | %s | %s | %s | %s |" % (
                title, size if size is not None else "—", share_url,
                fetched_at, summary.replace("|", "\\|"), COPYRIGHT_NOTE, round_ts))

        total_ok += n_ok
        total_fail += n_fail
        head = [f"### 链接{mi}：{share_url}", "",
                f"- 清单时间：{fetched_at}（第 {rnd} 轮）；目录 {len(dirs)} 个 / 文件 {len(files)} 个；"
                f"已获取正文 {n_ok} / 未获取正文 {n_fail}", "",
                "| 文件 | 大小(字节) | 来源链接 | 获取时间 | 内容摘要 | 版权与使用注记 | 轮次时间戳 |",
                "|---|---|---|---|---|---|---|"]
        sections.append("\n".join(head + rows))

    lines.append(f"- 总览：{len(manifest_paths)} 个分享链接，共 {total_files} 个文件；"
                 f"已获取正文 {total_ok} 个，未获取正文 {total_fail} 个"
                 f"（如实登记，未伪造任何下载成功或内容摘要）")
    lines.append("")
    lines.append("## 版权与使用总注记")
    lines.append("")
    lines.append(f"> {COPYRIGHT_NOTE}。若任何权利人主张权利，本登记册所载条目将配合删除。")
    lines.append("")
    lines.append("## 逐文件登记")
    lines.append("")
    lines.extend(sections)
    lines.append("")

    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return total_files, total_ok, total_fail


# ---------------------------------------------------------------- 冒烟自测
def smoke():
    """合成清单 + 合成证据：覆盖已下载(txt/pdf 摘要抽取)与 403 未获正文双路径。"""
    with tempfile.TemporaryDirectory(prefix="register_smoke_") as tmp:
        ldir = os.path.join(tmp, "SmokeHash123")
        os.makedirs(ldir)
        # 造一个已下载的 txt 与一个已下载的 pdf（pymupdf 缺失则降级跳过 pdf 路径）
        with open(os.path.join(ldir, "说明.txt"), "w", encoding="utf-8") as f:
            f.write("本资料为招生录取数据合辑，供内部参考。" * 10)
        pdf_made = False
        try:
            import fitz
            doc = fitz.open()
            page = doc.new_page()
            page.insert_text((72, 72), "2026 年硕士招生简章：报考条件、考试科目与录取办法。",
                             fontname="china-s")  # 内置 CJK 字体，保证中文可抽取
            doc.save(os.path.join(ldir, "招生简章.pdf"))
            doc.close()
            pdf_made = True
        except ImportError:
            pass

        share = "https://www.jianguoyun.com/p/SmokeHash123"
        objects = [
            {"name": "dir1", "relPath": "/dir1", "type": "directory",
             "size": 0, "mtime": "2026-08-12"},
            {"name": "说明.txt", "relPath": "/说明.txt", "type": "file",
             "size": 100, "mtime": "2026-08-12"},
            {"name": "录取数据.pdf", "relPath": "/dir1/录取数据.pdf", "type": "file",
             "size": 2048, "mtime": "2026-08-12"},
        ]
        if pdf_made:
            objects.append({"name": "招生简章.pdf", "relPath": "/招生简章.pdf",
                            "type": "file", "size": 3000, "mtime": "2026-08-12"})
        manifest = os.path.join(tmp, "清单.json")
        with open(manifest, "w", encoding="utf-8") as f:
            json.dump({"source": share, "fetched_at": "2026-08-24 05:57:00 +0800",
                       "round": 1, "objects": objects}, f, ensure_ascii=False)

        attempts = [
            {"link": "SmokeHash123", "relPath": "/说明.txt", "http_status": 200,
             "saved_path": "说明.txt", "saved_bytes": 100, "errorCode": None,
             "detailMsg": "OK"},
            {"link": "SmokeHash123", "relPath": "/dir1/录取数据.pdf",
             "http_status": 403, "saved_path": None, "saved_bytes": None,
             "errorCode": "SandboxAccessDenied", "detailMsg": "Can not download it"},
        ]
        if pdf_made:
            attempts.append({"link": "SmokeHash123", "relPath": "/招生简章.pdf",
                             "http_status": 200, "saved_path": "招生简章.pdf",
                             "saved_bytes": 3000, "errorCode": None, "detailMsg": "OK"})
        evidence = os.path.join(tmp, "下载尝试证据.json")
        with open(evidence, "w", encoding="utf-8") as f:
            json.dump({"round1": {"fetched_at": "2026-08-24 05:57:00 +0800",
                                  "download_attempts": attempts},
                       "round2": {"fetched_at": "2026-08-24 07:12:00 +0800",
                                  "download_attempts": attempts}},
                      f, ensure_ascii=False)

        out = os.path.join(tmp, "出处登记册.md")
        total, n_ok, n_fail = build_register([manifest], evidence, tmp, out)
        text = open(out, encoding="utf-8").read()

        assert COPYRIGHT_NOTE in text, "缺少字幕组式版权注记"
        assert share in text, "缺少来源链接"
        assert "招生录取数据合辑" in text, "txt 摘要未抽取"
        assert "未获取正文（分享者设置了“仅预览/禁止下载”权限，下载接口返回 403 " \
               "SandboxAccessDenied；本管线不绕过该限制）" in text, "403 诚实标注缺失"
        assert "round1 2026-08-24 05:57:00 +0800" in text, "轮次时间戳缺失"
        assert "round2" in text, "第 2 轮时间戳缺失"
        if pdf_made:
            assert "报考条件" in text, "pdf 摘要未抽取"
        else:
            print("[SMOKE NOTE] pymupdf 缺失，PDF 摘要路径未验证（降级）")
        assert total == len(objects) - 1 and n_fail == 1

        # 百度网盘清单兼容路径：source_url 键（baidu_share_browse.py 产物）
        baidu_manifest = os.path.join(tmp, "百度清单.json")
        with open(baidu_manifest, "w", encoding="utf-8") as f:
            json.dump({"source_url": "https://pan.baidu.com/s/1AbCdef",
                       "fetched_at": "2026-08-24 09:00:00 +0800", "round": 3,
                       "license_note": "仅限用户本人文件或获授权分享",
                       "summary": {"dirs": 1, "files": 1},
                       "objects": [{"name": "b.docx", "relPath": "/资料/b.docx",
                                    "type": "file", "size": 2048, "fsid": 21,
                                    "mtime": "2026-08-20 10:00:00 +0800",
                                    "category": 4, "category_name": "文档"}]},
                      f, ensure_ascii=False)
        out2 = os.path.join(tmp, "百度登记册.md")
        total2, _, n_fail2 = build_register([baidu_manifest], None, None, out2)
        text2 = open(out2, encoding="utf-8").read()
        assert total2 == 1 and n_fail2 == 1
        assert "https://pan.baidu.com/s/1AbCdef" in text2, "source_url 未进入登记册"
        assert "未获取正文（未在证据中找到下载记录" in text2, "无证据诚实标注缺失"

    print("[SMOKE OK] source_register: 清单+证据→登记册，摘要抽取/403标注/轮次时间戳/百度source_url兼容 全部通过")
    return 0


def main():
    ap = argparse.ArgumentParser(description="出处登记册生成器")
    ap.add_argument("--manifest", action="append", default=[],
                    help="文件清单 JSON（可重复传入多个）")
    ap.add_argument("--evidence", default=None, help="下载尝试证据 JSON")
    ap.add_argument("--archive-root", default=None, help="归档根目录（就地抽取摘要）")
    ap.add_argument("--out", default="出处登记册.md", help="输出路径")
    ap.add_argument("--smoke", action="store_true", help="合成清单自测，不联网")
    args = ap.parse_args()
    if args.smoke:
        return smoke()
    if not args.manifest:
        ap.error("至少传入一个 --manifest；或使用 --smoke 自测")
    total, n_ok, n_fail = build_register(
        args.manifest, args.evidence, args.archive_root, args.out)
    print(f"登记完成：{total} 个文件（已获正文 {n_ok} / 未获正文 {n_fail}）-> {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
